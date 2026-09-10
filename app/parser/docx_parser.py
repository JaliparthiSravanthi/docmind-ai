
from docx import Document
import os

# File paths
docx_path = "data/input/sample.docx"
output_path = "data/output/docx_text.txt"


def extract_text_from_docx(file_path):
    document = Document(file_path)
    full_text = ""

    print(f"Total Paragraphs: {len(document.paragraphs)}")

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()

        if text:
            full_text += text + "\n"
            print(f"Paragraph {index + 1}: {text}")

    return full_text


os.makedirs("data/output", exist_ok=True)

text = extract_text_from_docx(docx_path)

with open(output_path, "w", encoding="utf-8") as file:
    file.write(text)

print(f"\nText saved successfully to: {output_path}")